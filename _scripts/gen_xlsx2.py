# -*- coding: utf-8 -*-
"""Planilha da rede única (todos os planos somados), uma aba por cidade."""
import os, sys, json, re, shutil
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')
AQUI = os.path.dirname(os.path.abspath(__file__))
BASE = r'C:\Users\Dalla\OneDrive\Desktop\NOSSA SAÚDE\ESTUDO CAMPOS GERAIS'
P01 = os.path.join(BASE, '01 - REDE CREDENCIADA', 'Nossa Saúde')
os.makedirs(P01, exist_ok=True)

D = json.load(open(os.path.join(AQUI, 'dados_unico.json'), encoding='utf-8'))
ITENS, CIDADES, EXAMES = D['itens'], D['cidades'], D['exames']

NAVY = 'FF0D2A4F'
OURO = 'FF9A7513'
CINZA = 'FF89909C'
ALT = 'FFF7F8FB'
BORDA = 'FFE2E6EC'
CATS = ['Hospitais e Pronto-Socorro', 'Clínicas e Centros Médicos',
        'Laboratórios e Imagem', 'Profissionais (médicos e demais)']
CAT_CURTA = {CATS[0]: 'Hospitais', CATS[1]: 'Clínicas', CATS[2]: 'Labs e Imagem',
             CATS[3]: 'Profissionais'}
CAT_COR = {CATS[0]: 'FFB3261E', CATS[1]: 'FF1D4E89', CATS[2]: 'FF1F7A5A', CATS[3]: 'FF6B4E9B'}

F_TIT = Font(size=16, bold=True, color=NAVY)
F_SUB = Font(size=10, color=CINZA)
F_HDR = Font(size=10, bold=True, color='FFFFFFFF')
F_TXT = Font(size=10)
F_B = Font(size=10, bold=True)
FILL_HDR = PatternFill('solid', fgColor=NAVY)
FILL_ALT = PatternFill('solid', fgColor=ALT)
FILL_OURO = PatternFill('solid', fgColor=OURO)
THIN = Side(style='thin', color=BORDA)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

AVISO = ('A rede credenciada é definida e alterada exclusivamente pela operadora. Esta lista soma '
         'todos os planos da Nossa Saúde — confirme no portal da operadora se o prestador atende o '
         'plano específico do cliente antes de contratar.')
FONTE = ('Fonte: Rede Credenciada oficial da Nossa Saúde (prestador.nossasaude.com.br), '
         'consulta em 10/09/2026, sem filtro de plano.')

COLS = [('Cidade', 16), ('Categoria', 26), ('Tipo (operadora)', 24), ('Prestador', 44),
        ('Razão social / Nome', 44), ('CNPJ', 20), ('Conselho', 18), ('Especialidades', 46),
        ('Exames que realiza', 30), ('Internação', 12), ('Endereço', 50), ('Bairro', 20),
        ('CEP', 12), ('Telefones', 32), ('Corpo clínico', 60)]
WRAP = {8, 9, 11, 14, 15}


def tels(r):
    out = []
    for e in r['enderecos']:
        for t in re.split(r'\s*/\s*', e.get('tel') or ''):
            t = t.strip()
            if t and t not in out:
                out.append(t)
    return out


def ends(r):
    out = []
    for e in r['enderecos']:
        p = e['logradouro']
        if e['cep']:
            p += ' — CEP ' + e['cep']
        if p and p not in out:
            out.append(p)
    return out


def cab(ws, titulo, sub):
    ws['A1'] = titulo; ws['A1'].font = F_TIT
    ws['A2'] = sub; ws['A2'].font = F_SUB
    ws['A3'] = AVISO; ws['A3'].font = F_SUB
    ws.freeze_panes = 'A6'


def tabela(ws, recs, r0=5):
    for j, (h, w) in enumerate(COLS, start=1):
        c = ws.cell(row=r0, column=j, value=h)
        c.font = F_HDR; c.fill = FILL_HDR; c.border = BOX
        c.alignment = Alignment(vertical='center')
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[r0].height = 22
    r = r0 + 1
    for rec in recs:
        vals = [rec['cidade'], rec['categoria'], rec['tipo_exib'], rec['nome_exib'], rec['razao'],
                rec['cnpj'], rec['conselho'], ' • '.join(rec['esp_exib']),
                ' • '.join(rec['naturezas']), 'sim' if rec['internacao'] else '',
                ' | '.join(ends(rec)),
                ' | '.join(dict.fromkeys(e['bairro'] for e in rec['enderecos'] if e['bairro'])),
                ' | '.join(dict.fromkeys(e['cep'] for e in rec['enderecos'] if e['cep'])),
                ' | '.join(tels(rec)),
                ' | '.join('%s (%s)' % (m['nome'], m['registro']) for m in rec['corpo_clinico'])]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=j, value=v)
            c.font = F_TXT; c.border = BOX
            c.alignment = Alignment(vertical='top', wrap_text=(j in WRAP))
            if r % 2 == 0:
                c.fill = FILL_ALT
            if j == 4:
                c.font = F_B
            elif j == 2:
                c.font = Font(size=10, bold=True, color=CAT_COR[rec['categoria']])
            elif j == 10 and v:
                c.font = Font(size=10, bold=True, color=OURO)
        r += 1
    ws.auto_filter.ref = 'A%d:%s%d' % (r0, get_column_letter(len(COLS)), r - 1)


wb = Workbook()

# ------------------------------------------------ RESUMO
ws = wb.active; ws.title = 'RESUMO'
ws['A1'] = 'REDE CREDENCIADA — NOSSA SAÚDE'; ws['A1'].font = Font(size=18, bold=True, color=NAVY)
ws['A2'] = 'Campos Gerais / PR · todos os planos somados numa rede só · Mazza Broker'
ws['A2'].font = Font(size=11, bold=True, color=OURO)
ws['A3'] = FONTE; ws['A3'].font = F_SUB
ws['A4'] = AVISO; ws['A4'].font = F_SUB
for w, col in zip([30, 14, 14, 16, 16, 14, 14], 'ABCDEFG'):
    ws.column_dimensions[col].width = w

r = 6
hdr = ['Cidade', 'Hospitais', 'Labs e imagem', 'Clínicas', 'Profissionais', 'Total']
for j, h in enumerate(hdr, start=1):
    c = ws.cell(row=r, column=j, value=h); c.font = F_HDR; c.fill = FILL_HDR; c.border = BOX
r += 1
tot = [0, 0, 0, 0]
for cid in CIDADES:
    lst = [i for i in ITENS if i['cidade'] == cid]
    n = [sum(1 for i in lst if i['internacao']),
         sum(1 for i in lst if i['naturezas'] and not i['internacao']),
         sum(1 for i in lst if i['categoria'] == CATS[1]),
         sum(1 for i in lst if i['categoria'] == CATS[3])]
    tot = [a + b for a, b in zip(tot, n)]
    for j, v in enumerate([cid] + n + [len(lst)], start=1):
        c = ws.cell(row=r, column=j, value=v); c.border = BOX
        c.font = F_B if j in (1, 6) else F_TXT
    r += 1
for j, v in enumerate(['TOTAL'] + tot + [len(ITENS)], start=1):
    c = ws.cell(row=r, column=j, value=v)
    c.font = Font(size=10, bold=True, color='FFFFFFFF'); c.fill = FILL_OURO; c.border = BOX
r += 3

ws.cell(row=r, column=1, value='EXAMES E DIAGNÓSTICO — QUANTOS PRESTADORES POR NATUREZA').font = Font(size=12, bold=True, color=NAVY)
r += 2
for j, h in enumerate(['Natureza do exame'] + CIDADES + ['Total'], start=1):
    c = ws.cell(row=r, column=j, value=h); c.font = F_HDR; c.fill = FILL_HDR; c.border = BOX
    ws.column_dimensions[get_column_letter(j)].width = max(14, ws.column_dimensions[get_column_letter(j)].width or 0)
r += 1
for nat in EXAMES:
    lst = [i for i in ITENS if nat in i['naturezas']]
    vals = [nat] + [sum(1 for i in lst if i['cidade'] == c) or '' for c in CIDADES] + [len(lst)]
    for j, v in enumerate(vals, start=1):
        c = ws.cell(row=r, column=j, value=v); c.border = BOX
        c.font = F_B if j == 1 else F_TXT
    r += 1
r += 2
ws.cell(row=r, column=1, value='Contato: Alan Vinicius Dall Alba · (41) 99547-6715 · '
                               'alan.vinicius@mazzabroker.com.br').font = F_SUB

# ------------------------------------------------ rede completa + cidades
ordem_cat = lambda x: (CATS.index(x['categoria']), x['nome_exib'].lower())
ws = wb.create_sheet('REDE COMPLETA')
todos = sorted(ITENS, key=lambda x: (CIDADES.index(x['cidade']),) + ordem_cat(x))
cab(ws, 'Rede credenciada Nossa Saúde — Campos Gerais',
    '%d prestadores nas 9 cidades · %s' % (len(todos), FONTE))
tabela(ws, todos)

for cid in CIDADES:
    ws = wb.create_sheet(cid[:31])
    recs = sorted([i for i in ITENS if i['cidade'] == cid], key=ordem_cat)
    cab(ws, '%s / PR' % cid, '%d prestadores · %s' % (len(recs), FONTE))
    tabela(ws, recs)

for velho in os.listdir(P01):
    if velho.endswith('.xlsx'):
        os.remove(os.path.join(P01, velho))
out = os.path.join(P01, 'REDE NOSSA SAUDE - CAMPOS GERAIS.xlsx')
wb.save(out)
print('OK', out, len(todos), 'prestadores')

site = os.path.join(AQUI, 'site', 'arquivos')
if os.path.isdir(site):
    wb.save(os.path.join(site, 'rede-nossa-saude-campos-gerais.xlsx'))
    print('   copiado para o site')
